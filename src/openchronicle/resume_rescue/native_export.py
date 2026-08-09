"""Deterministic native document exports from the reviewed résumé tree."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import PurePosixPath
from typing import Any

from defusedxml import ElementTree as DefusedElementTree
from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .render import RENDERER_VERSION, TEMPLATE_ID, ResumeDocumentTree

NATIVE_EXPORT_VERSION = 1
MAX_DOCX_BYTES = 4 * 1024 * 1024
MAX_DOCX_MEMBERS = 100
MAX_DOCX_UNCOMPRESSED_BYTES = 8 * 1024 * 1024
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
_REL = f"{{{_RELATIONSHIP_NAMESPACE}}}"
_ACTIVE_MARKERS = ("vbaproject", "activex/", "embeddings/", "oleobject")
_FIXED_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)


class NativeResumeExportError(ValueError):
    """A native résumé document could not be built or validated."""


@dataclass(frozen=True, slots=True)
class ResumeNativeExport:
    projection_id: str
    artifact_digest: str
    preview_document_digest: str
    format: str
    media_type: str
    extension: str
    content: bytes
    content_digest: str

    def metadata(self) -> dict[str, Any]:
        return {
            "schema_version": 1,
            "projection_id": self.projection_id,
            "artifact_digest": self.artifact_digest,
            "preview_document_digest": self.preview_document_digest,
            "renderer_version": RENDERER_VERSION,
            "native_export_version": NATIVE_EXPORT_VERSION,
            "template_id": TEMPLATE_ID,
            "format": self.format,
            "media_type": self.media_type,
            "extension": self.extension,
            "byte_count": len(self.content),
            "content_digest": self.content_digest,
            "action_capability": "none",
        }


def render_docx_export(
    tree: ResumeDocumentTree, *, preview_document_digest: str
) -> ResumeNativeExport:
    """Render and audit a deterministic DOCX from the shared semantic tree."""

    if not isinstance(tree, ResumeDocumentTree) or not re.fullmatch(
        r"[0-9a-f]{64}", preview_document_digest
    ):
        raise NativeResumeExportError("native export binding is invalid")
    raw = _build_docx(tree)
    content = _canonicalize_docx(raw)
    _validate_docx(content, tree)
    return ResumeNativeExport(
        projection_id=tree.projection_id,
        artifact_digest=tree.artifact_digest,
        preview_document_digest=preview_document_digest,
        format="docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        extension="docx",
        content=content,
        content_digest=hashlib.sha256(content).hexdigest(),
    )


def _build_docx(tree: ResumeDocumentTree) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(17)
    section.right_margin = Mm(17)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    name.paragraph_format.space_after = Pt(12)
    name.paragraph_format.keep_with_next = True
    name_run = name.add_run(tree.display_name)
    name_run.bold = True
    name_run.font.name = "Arial"
    name_run.font.size = Pt(25)
    name_run.font.color.rgb = RGBColor(23, 32, 29)
    name_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")

    for semantic_section in tree.sections:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(5)
        heading.paragraph_format.keep_with_next = True
        heading_run = heading.add_run(semantic_section.label.upper())
        heading_run.bold = True
        heading_run.font.name = "Arial"
        heading_run.font.size = Pt(11)
        heading_run.font.color.rgb = RGBColor(23, 107, 92)
        heading_run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
        for item in semantic_section.items:
            paragraph = document.add_paragraph(style="List Bullet")
            _set_explicit_bullet_numbering(paragraph)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(item.text)
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")

    properties = document.core_properties
    properties.title = f"{tree.display_name} — Résumé"
    properties.subject = "Reviewed OpenChronicle Résumé Rescue export"
    properties.author = "OpenChronicle"
    properties.last_modified_by = "OpenChronicle"
    properties.keywords = (
        f"projection={tree.projection_id};artifact={tree.artifact_digest};"
        f"template={TEMPLATE_ID};renderer={RENDERER_VERSION}"
    )
    timestamp = _stable_datetime(tree.created_at)
    properties.created = timestamp
    properties.modified = timestamp
    _normalize_bullet_numbering(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _set_explicit_bullet_numbering(paragraph: Paragraph) -> None:
    """Keep bullets visible when office suites ignore style-only numbering."""

    properties = paragraph._p.get_or_add_pPr()
    numbering = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), "1")
    numbering.append(level)
    numbering.append(number)
    properties.append(numbering)


def _normalize_bullet_numbering(document: DocxDocument) -> None:
    """Use a Unicode bullet instead of the template's private-use Symbol glyph."""

    numbering = document.part.numbering_part.element
    for abstract in numbering.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            number_format = level.find(qn("w:numFmt"))
            level_text = level.find(qn("w:lvlText"))
            if (
                number_format is None
                or number_format.get(qn("w:val")) != "bullet"
                or level_text is None
            ):
                continue
            level_text.set(qn("w:val"), "•")
            run_properties = level.find(qn("w:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("w:rPr")
                level.append(run_properties)
            fonts = run_properties.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                run_properties.append(fonts)
            for attribute in ("ascii", "hAnsi", "eastAsia"):
                fonts.set(qn(f"w:{attribute}"), "Arial")


def _stable_datetime(value: str) -> datetime:
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise NativeResumeExportError("native export timestamp is invalid") from exc
    if parsed.tzinfo is None:
        raise NativeResumeExportError("native export timestamp is invalid")
    return parsed.astimezone(UTC).replace(tzinfo=None, microsecond=0)


def _canonicalize_docx(source: bytes) -> bytes:
    output = io.BytesIO()
    try:
        with zipfile.ZipFile(io.BytesIO(source)) as package:
            names = package.namelist()
            if len(names) != len(set(names)):
                raise NativeResumeExportError("DOCX contains duplicate members")
            values = [(name, package.read(name)) for name in sorted(names)]
        with zipfile.ZipFile(
            output,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
            strict_timestamps=True,
        ) as package:
            for name, value in values:
                info = zipfile.ZipInfo(name, _FIXED_ZIP_TIMESTAMP)
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = 0o100600 << 16
                package.writestr(info, value, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    except (OSError, zipfile.BadZipFile) as exc:
        raise NativeResumeExportError("DOCX package could not be canonicalized") from exc
    return output.getvalue()


def _validate_docx(source: bytes, tree: ResumeDocumentTree) -> None:
    if not source.startswith(b"PK") or not 1 <= len(source) <= MAX_DOCX_BYTES:
        raise NativeResumeExportError("DOCX output size or signature is invalid")
    try:
        with zipfile.ZipFile(io.BytesIO(source)) as package:
            infos = package.infolist()
            if not infos or len(infos) > MAX_DOCX_MEMBERS:
                raise NativeResumeExportError("DOCX member count is invalid")
            if [info.filename for info in infos] != sorted(info.filename for info in infos):
                raise NativeResumeExportError("DOCX member order is not canonical")
            total = 0
            for info in infos:
                path = PurePosixPath(info.filename)
                if (
                    not info.filename
                    or "\\" in info.filename
                    or path.is_absolute()
                    or ".." in path.parts
                    or info.flag_bits & 0x1
                    or any(marker in info.filename.casefold() for marker in _ACTIVE_MARKERS)
                    or info.date_time != _FIXED_ZIP_TIMESTAMP
                ):
                    raise NativeResumeExportError("DOCX member is unsafe")
                total += info.file_size
                if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise NativeResumeExportError("DOCX expanded size exceeds limit")
                value = package.read(info)
                if info.filename.endswith((".xml", ".rels")):
                    root = DefusedElementTree.fromstring(value)
                    if info.filename.endswith(".rels"):
                        for relationship in root.findall(f"{_REL}Relationship"):
                            if relationship.attrib.get("TargetMode") == "External":
                                raise NativeResumeExportError(
                                    "DOCX external relationships are forbidden"
                                )
            names = {info.filename for info in infos}
            if not {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= names:
                raise NativeResumeExportError("DOCX required members are missing")
    except (DefusedElementTree.ParseError, KeyError, zipfile.BadZipFile) as exc:
        raise NativeResumeExportError("DOCX output could not be validated") from exc

    parsed = Document(io.BytesIO(source))
    actual = [paragraph.text for paragraph in parsed.paragraphs if paragraph.text]
    expected = [tree.display_name]
    for section in tree.sections:
        expected.append(section.label.upper())
        expected.extend(item.text for item in section.items)
    if actual != expected:
        raise NativeResumeExportError("DOCX ordered text differs from the semantic tree")
